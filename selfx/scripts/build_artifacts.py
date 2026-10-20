"""Build figures, editable report and slide content from saved experiment files.

Run with a Python environment containing requirements-dev.txt. PPTX authoring is
handled by build_presentation.mjs using the Codex Artifact Tool/session.
"""
from __future__ import annotations
import csv
import json
import math
from pathlib import Path
from statistics import mean

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
FINAL = OUT / "submission"
FIG = OUT / "figures"
COLORS = ["#a4b4b1", "#6e9697", "#227d87", "#007d64"]
SOURCES = [
    ("Shinn et al. (2023). Reflexion: Language Agents with Verbal Reinforcement Learning. NeurIPS.", "https://arxiv.org/abs/2303.11366"),
    ("Rusek et al. (2019). RouteNet: Leveraging Graph Neural Networks for network modeling and optimization in SDN.", "https://arxiv.org/abs/1910.01508"),
    ("Google. Gemini API model catalog.", "https://ai.google.dev/gemini-api/docs/models"),
    ("Google. Structured outputs.", "https://ai.google.dev/gemini-api/docs/structured-output"),
    ("Google. Gemini API rate limits.", "https://ai.google.dev/gemini-api/docs/rate-limits"),
    ("OpenRouter. Free Models Router.", "https://openrouter.ai/docs/guides/routing/routers/free-router"),
    ("Ryu. ofctl_rest reference.", "https://ryu.readthedocs.io/en/latest/app/ofctl_rest.html"),
]


def load_run(name):
    folder = OUT / "experiments" / name
    return {"summary": json.loads((folder / "summary.json").read_text(encoding="utf-8")),
            "rows": list(csv.DictReader((folder / "metrics.csv").open(encoding="utf-8"))),
            "audit": json.loads((folder / "provider_calls.json").read_text(encoding="utf-8"))}


def coverage(run, condition):
    rows = [r for r in run["rows"] if r["condition"] == condition]
    return sum(r["decision_origin"] == "api" for r in rows), len(rows)


def figures(runs):
    FIG.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False,
                         "axes.spines.right": False, "axes.labelcolor": "#17343a", "text.color": "#17343a"})
    summary = runs["offline"]["summary"]["conditions"]
    fig, axes = plt.subplots(1, 3, figsize=(11.8, 3.45), constrained_layout=True)
    for ax, key, title, scale in zip(axes, ["latency_ms", "throughput_kbps", "loss_pct"],
                                    ["Modeled latency (ms)", "Delivered goodput (Mbps)", "Mean total loss (%)"], [1, .001, 1]):
        values = [summary[c][key] * scale for c in summary]
        bars = ax.bar(list(summary), values, color=COLORS, width=.6)
        ax.set_title(title, loc="left", fontsize=11, pad=15)
        ax.set_ylim(0, max(values) * 1.23)
        ax.bar_label(bars, labels=[f"{v:.2f}" for v in values], padding=4, fontsize=9)
        ax.grid(axis="y", alpha=.15)
        ax.set_axisbelow(True)
    fig.savefig(FIG / "offline_comparison.png", dpi=210, facecolor="white")
    plt.close(fig)
    fig, axes = plt.subplots(1, 2, figsize=(11.8, 3.8), constrained_layout=True)
    for ax, run_name in zip(axes, ["gemini_paced", "openrouter"]):
        run = runs[run_name]
        s = run["summary"]["conditions"]
        vals = [s[c]["latency_ms"] for c in s]
        bars = ax.bar(list(s), vals, color=COLORS)
        ax.bar_label(bars, labels=[f"{v:.2f}" for v in vals], padding=4, fontsize=9)
        ax.set_title(("Gemini paced" if run_name == "gemini_paced" else "OpenRouter router diagnostic") + " | modeled latency (ms)", loc="left", fontsize=10, pad=14)
        ax.set_ylim(0, max(vals)*1.25)
        ax.grid(axis="y", alpha=.15); ax.set_axisbelow(True)
    fig.savefig(FIG / "api_latency.png", dpi=210, facecolor="white")
    plt.close(fig)
    fig, axes = plt.subplots(1, 2, figsize=(11.8, 3.5), constrained_layout=True)
    for ax, run_name in zip(axes, ["gemini_paced", "openrouter"]):
        run=runs[run_name]
        counts=[coverage(run,c) for c in ("A1","A2")]
        ok=[x[0] for x in counts]; fallback=[x[1]-x[0] for x in counts]
        ax.bar(["A1","A2"],ok,color="#007d64",label="API decision")
        ax.bar(["A1","A2"],fallback,bottom=ok,color="#d8a46a",label="safe fallback")
        ax.set_title(run_name.replace("_"," "),loc="left")
        ax.set_ylabel("Control windows")
        for i,(a,b) in enumerate(counts): ax.text(i,b+1,f"{a}/{b} API",ha="center",fontsize=10)
        ax.set_ylim(0,max(x[1] for x in counts)*1.2)
        ax.legend(frameon=False,loc="upper right",fontsize=8)
    fig.savefig(FIG / "api_coverage.png", dpi=210, facecolor="white");plt.close(fig)
    rows=runs["offline"]["rows"]
    fig,ax=plt.subplots(figsize=(11.8,3.2),constrained_layout=True)
    for condition,color in zip(("A1","A2"),COLORS[2:]):
        vals=[]
        for cycle in range(120):
            vals.append(mean(float(r["reward"]) for r in rows if r["condition"]==condition and int(r["cycle"])==cycle))
        rolling=[mean(vals[max(0,i-9):i+1]) for i in range(len(vals))]
        ax.plot(range(1,121),rolling,label=condition,color=color,linewidth=2)
    ax.axhline(0,color="#a4b4b1",linewidth=.7);ax.set_xlabel("Window index");ax.set_ylabel("10-window rolling mean reward")
    ax.legend(frameon=False);ax.grid(alpha=.12)
    fig.savefig(FIG / "offline_reward.png",dpi=210,facecolor="white");plt.close(fig)


def style_document(doc):
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.72)
    sec.left_margin=sec.right_margin=Inches(.78)
    for name in ("Normal","Title","Subtitle","Heading 1","Heading 2","Caption"):
        st=doc.styles[name];st.font.name="Arial";st.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles["Normal"];normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
    for name,size in (("Title",29),("Heading 1",20),("Heading 2",12.5),("Subtitle",12)):
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].paragraph_format.space_after=Pt(10)
    foot=sec.footer.paragraphs[0];foot.alignment=2
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");foot._p.append(field)
    doc.core_properties.title="SELF X SDN control with state embedding memory"
    doc.core_properties.author="SELF-X Project"
    doc.core_properties.subject="Measured simulation and live provider integration"


def para(doc,text):
    doc.add_paragraph(text)


def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";t.autofit=False
    for i,h in enumerate(headers): t.rows[0].cells[i].text=h
    for data in rows:
        cells=t.add_row().cells
        for i,value in enumerate(data): cells[i].text=str(value)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            if widths: cell.width=Inches(widths[ci])
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
                for run in p.runs: run.font.size=Pt(9);run.bold=ri==0
            if ri==0:
                shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"EDF2EF");cell._tc.get_or_add_tcPr().append(shade)
        props=row._tr.get_or_add_trPr();cant=OxmlElement("w:cantSplit");props.append(cant)
    repeat=OxmlElement("w:tblHeader");t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t


def page(doc,title):
    doc.add_page_break();doc.add_heading(title,level=1)


def image(doc,name,caption):
    doc.add_picture(str(FIG/name),width=Inches(6.65))
    doc.add_paragraph(caption,style="Caption")


def result_table(doc,run):
    c=run["summary"]["conditions"]
    table(doc,["Controller","Latency ms","Goodput Mbps","Mean loss %","Vetoes"],
          [[k,f"{v['latency_ms']:.3f}",f"{v['throughput_kbps']/1000:.3f}",f"{v['loss_pct']:.3f}",v["safety_vetoes"]] for k,v in c.items()],
          [1.1,1.3,1.55,1.4,1.15])


def make_report(runs,checks):
    offline=runs["offline"]["summary"];c=offline["conditions"]
    gemini=runs["gemini_paced"];gc=gemini["summary"]["conditions"]
    orouter=runs["openrouter"]
    total=len(checks["models"]);passed=sum(r["status"]=="ok" for r in checks["models"])
    delta=c["A2"]["latency_ms"]-c["A1"]["latency_ms"]
    doc=Document();style_document(doc)
    doc.add_paragraph("SELF X SDN control with state embedding memory",style="Title")
    doc.add_paragraph("Project report | 20 October 2026",style="Subtitle")
    doc.add_heading("Abstract",level=1)
    para(doc,"SELF-X investigates a single software-defined networking controller that retrieves earlier state-action-outcome episodes and their reflections before choosing a flow intervention. The project combines typed JSON decisions, a deterministic action gate, a normalized numeric state embedding, bounded episode memory and measurable feedback. A local website exposes the complete decision trace and provides reproducible experiments.")
    para(doc,"The evaluation includes a finite-buffer two-path fluid simulation with 2,400 windows across five seeds and four controllers, a paced Gemini 3.5 Flash-Lite pilot, and an OpenRouter free-model integration diagnostic. API calls are real; all network performance values in this report are simulated. A separate Linux Mininet/Ryu adapter is retained but has not been validated as a physical or packet-emulation testbed in this Windows evaluation.")
    para(doc,f"In the offline study, A1 mean modeled latency was {c['A1']['latency_ms']:.2f} ms and A2 was {c['A2']['latency_ms']:.2f} ms. Memory therefore changed latency by {delta:+.2f} ms relative to the same controller without memory. A2 also had {c['A2']['loss_pct']:.2f}% mean total loss compared with {c['A1']['loss_pct']:.2f}% for A1. These results do not establish a consistent benefit from memory. They demonstrate why safety, actual API coverage and multiple metrics must accompany the learning claim.")
    if (OUT/'benchmark_chart_data.json').exists():
        para(doc,"The expanded evaluation adds 57,600 windows across six traffic scenarios, eight controllers and ten paired seeds. Guarded balanced memory A3 had no rejected proposals, compared with 134 for original memory A2. Against the same guarded controller without memory H1, A3 increased mixed-load delay by 0.473 ms; memory did not show a consistent performance gain. Appendices A-F document uncertainty, tradeoffs and the isolated virtual-environment execution.")
    doc.add_heading("Report navigation",level=2)
    para(doc,"The main sections cover the objective, related work, architecture, method, original experiments, API pilots, validation, limitations and references. Appendices A-F cover the expanded benchmark, guard ablation and fresh API run. Raw evidence is under output/experiments and output/benchmarks/v2; runnable instructions are in README.md.")

    page(doc,"1 Problem and project objective")
    para(doc,"Network demand varies over time. A long-lived or bursty high-bandwidth flow can congest a shared bottleneck, increase queueing delay and trigger loss. A static route may continue to use the overloaded path while an alternate path has spare capacity. An administrator needs both an actionable control choice and enough evidence to assess its consequences.")
    para(doc,"The research question is whether a single controller can improve decisions during repeated traffic episodes by retrieving its own previous experiences. Each experience records a compact state, the installed action, measured outcome, reward and reflection. The later controller can consult similar episodes without retraining its language model.")
    doc.add_heading("Research questions",level=2)
    for text in ["Can the implementation observe traffic, request a structured decision, reject invalid actions and finish an auditable control cycle?", "Does A2 improve latency, delivered goodput or total loss over the otherwise comparable A1 controller?", "How do quota errors, malformed responses and action vetoes affect operational results?", "Can a reviewer reproduce the traffic and trace a reported metric back to a specific window?"]:
        para(doc,text)
    doc.add_heading("Intended website behavior",level=2)
    para(doc,"The website is a local research dashboard. It compares controllers, replays an offline experiment from a seed, and exposes the exact decision, safety reason and reflection for each window. It also links the editable report and presentation. Provider credentials remain in Python and are never served to the browser.")
    para(doc,"The submission focuses only on the selfx folder. Its portable model is designed for a reliable demonstration on this Windows computer. Deployment on a real SDN requires separate testbed instrumentation and validation; a successful local simulation is not evidence of deployment performance.")

    page(doc,"2 Related work and comparison boundary")
    para(doc,"Reflexion uses verbal feedback and episodic memory to guide language agents over repeated attempts [1]. It motivates storing useful explanations beside measured feedback without fine-tuning. Its published evaluation tasks differ from network flow management, so its task scores cannot be interpreted as an expected SDN improvement.")
    para(doc,"RouteNet models relationships among topology, routing and traffic using a graph neural network [2]. The paper reports worst-case mean relative error of 15.4% in its KPI prediction evaluation. This is prediction error under that study's data and split, not a latency reduction and not a baseline reproduced by SELF-X. Its results support the relevance of topology-aware modeling but require a separate implementation and dataset for a fair numerical comparison.")
    table(doc,["Approach","Role in this report","Comparable numerical result"],[
        ["Static primary path A0","Local no-control baseline","Measured on identical traffic"],
        ["Greedy route B1","Local deterministic baseline","Measured on identical traffic"],
        ["Reflexion","Conceptual basis for episode reflections","No cross-task score comparison"],
        ["RouteNet","Published learned network KPI model","15.4% MRE belongs to original study"],
        ["Gemini and OpenRouter","Live decision providers","Coverage and local simulated outcome"],
    ],[1.35,2.4,2.75])
    para(doc,"The local baselines require no model training and make no claim to reproduce a particular paper's routing algorithm. B1 moves the elephant flow to the less-loaded path only when the alternate remains within the safety ceiling. It provides a simple operational comparator. Offline A1/A2 evaluate deterministic threshold and case-retrieval policies; API pilots evaluate actual provider integration separately.")
    para(doc,"Official Gemini documentation identifies the requested model endpoint and explains structured output and account-specific rate limits [3-5]. OpenRouter documents its free router and routed model selection [6]. Ryu documents the flow and meter REST interface used by the separate Linux adapter [7]. All external sources were accessed on 20 October 2026.")

    page(doc,"3 System architecture")
    para(doc,"The portable controller follows an explicit observe-retrieve-decide-validate-act-measure-reflect-store loop. Shared schemas, prompts, memory, reward and safety modules support the portable runner and original Ryu-oriented LangGraph pipeline. The simulator orchestrates these modules directly; the original testbed uses LangGraph nodes.")
    table(doc,["Stage","Input","Output"],[
        ["Observe","Seeded demand and current queues","State summary and numeric features"],
        ["Retrieve","Current embedding and earlier A2 records","Up to four similar episodes"],
        ["Decide","State plus examples","Typed action and rationale"],
        ["Validate and act","Decision and known topology","Allowed intervention or visible no-op"],
        ["Measure","Demand and selected action","Delay goodput loss and reward"],
        ["Reflect and store","Outcome and reward","Explanation plus bounded episode"],
    ],[1.35,2.3,2.85])
    doc.add_heading("Topology and action lifetime",level=2)
    para(doc,"The portable topology represents primary s1-s2-s4 and alternate s1-s3-s4 logical paths. Each path has a 10 Mbps bottleneck, a 0.5 Mbit queue and an 8 ms propagation component. Background traffic occupies each path independently. One elephant flow starts on the primary in each five-second window. An intervention applies only to that window; queues carry forward within the condition and seed.")
    para(doc,"The original Mininet topology is linear and has no alternate route. The Ryu graph therefore requires an explicitly validated egress mapping before a reroute is allowed. The portable diamond is a simulation abstraction, not eight instantiated hosts or a deployed OpenFlow topology.")
    doc.add_heading("Navigation and data flow",level=2)
    para(doc,"README.md is the entry point. docs/ARCHITECTURE.md and its R Markdown companion map every module and contain flow diagrams. The project guide explains the purpose and demo. Experiment files connect fixed traffic hashes to CSV measurements, JSONL decisions and provider audits. The report and charts are generated from those saved records.")

    page(doc,"4 Embedding memory and deterministic safety")
    para(doc,"The state embedding contains normalized primary background, alternate background, elephant demand, both queue occupancies and total primary demand, plus a small constant feature. Each vector is normalized to unit length. Cosine similarity ranks retained episodes. This is a hand-designed numeric representation; no neural embedding model or vector database is trained or queried.")
    para(doc,"A2 retrieves at most four earlier episodes. Each record contains state, installed action, current parameters, reward and reflection. Memory retains at most 30 records and prunes the lowest rewards. A1 disables record persistence and retrieves no examples. Every seed and condition starts empty, preventing future-window and cross-run leakage.")
    para(doc,"Offline A1 uses an 85% primary-load trigger. If the alternate can accept the whole elephant flow without exceeding 90% utilization, it reroutes; otherwise it may cap the elephant at 80% of offered demand. Offline A2 can reuse a positively rewarded action from a state with cosine similarity above 0.97, rebuilding parameters for current demand. The independent safety gate still runs. API A1/A2 instead request provider decisions with the same prompt apart from memory examples.")
    table(doc,["Check","Enforced behavior"],[
        ["Action schema","Only reroute rate_limit or no_op with required parameters"],
        ["Target validity","Exact known switch flow match and alternate egress"],
        ["Rate reduction","At most 20% of the current elephant demand"],
        ["Meter floor","At least 512 kbps"],
        ["Alternate capacity","Projected offered load no more than 90% of 10 Mbps"],
        ["Failure handling","Record the reason and apply no_op"],
    ],[1.55,4.95])
    para(doc,"These local constraints bound specific interventions. They do not prove global stability, fairness, SLA compliance or end-to-end loss limits. Similarity can retrieve an action inappropriate for the present state, and reward-based pruning can discard useful negative examples. The experiment records vetoes to expose those weaknesses.")

    page(doc,"5 Experimental protocol and network model")
    table(doc,["Setting","Main portable run","Gemini pilot","Free-router diagnostic"],[
        ["Seeds","5","3","1"],["Windows per seed and condition","120","24","12"],
        ["Simulated seconds each","600","120","60"],["Controllers","A0 B1 A1 A2","A0 B1 A1 A2","A0 B1 A1 A2"],
        ["Network provenance","Simulation","Simulation","Simulation"],
        ["A1 A2 decisions","Deterministic","Real Gemini or marked fallback","Real free router or marked fallback"],
    ],[2.0,1.4,1.4,1.7])
    para(doc,"The workload repeats light, primary-burst, alternate-busy and dual-congestion episodes with seeded uniform demand perturbations. It is synthetic aggregate NetFlow-style data rather than a captured traffic dataset. Every condition uses identical demand bytes and starts with zero queues and empty memory. SHA256 hashes in each summary bind the workload to its results.")
    doc.add_heading("Fluid equations",level=2)
    para(doc,"For each path, q_next = min(B, max(0, q_previous + (admitted_rate - C)T)). Overflow is max(0, q_previous + (admitted_rate - C)T - B). Delivered volume is old queue plus admitted volume minus the new queue and overflow. C is 10 Mbps, B is 0.5 Mbit and T is 5 seconds.")
    para(doc,"The exact average fluid queue is integrated over the window, including saturation or emptying. Delay equals 8 ms plus average queue/C converted to milliseconds, weighted by delivered volume across paths. Goodput is delivered volume/T. Loss includes overflow and deliberate meter drops against all offered volume. The result is a modeled delay proxy, not measured packet RTT, jitter or p95.")
    doc.add_heading("Reward and uncertainty",level=2)
    para(doc,"Reward = -0.4(delta latency/50) + 0.4(delta goodput/5000) - 0.2(delta loss/20), using ms, kbps and percentage points. A same-window no-op counterfactual uses the same starting queue and demand; only the chosen action commits the next queue. Paired seed means, rather than individual windows, define the 95% Student-t intervals. With one seed no interval is reported. Small-sample intervals are exploratory.")

    page(doc,"6 Main offline results")
    result_table(doc,runs["offline"])
    image(doc,"offline_comparison.png","Figure 1. Five-seed offline simulation. All mean loss values include deliberate meter drops. Each controller has 600 evaluated windows.")
    para(doc,f"A2 reduced mean modeled latency relative to A0 from {c['A0']['latency_ms']:.2f} to {c['A2']['latency_ms']:.2f} ms, but A1 was lower at {c['A1']['latency_ms']:.2f} ms. A2 had {c['A2']['safety_vetoes']} vetoes; similar-state reuse sometimes proposed an action that was not feasible under current demand. The safety gate converted these actions to no-op.")
    para(doc,f"A2 goodput was {c['A2']['throughput_kbps']/1000:.3f} Mbps, compared with {c['A1']['throughput_kbps']/1000:.3f} Mbps for A1. Mean total loss rose from {c['A1']['loss_pct']:.3f}% to {c['A2']['loss_pct']:.3f}%. A lower delay alone therefore cannot establish a better overall controller. B1 achieved the lowest total loss among these four means while A1 achieved the lowest delay.")
    rows=[]
    for condition in ("B1","A1","A2"):
        value=offline["paired_seed_comparisons"][condition+"_minus_A0"]["latency_ms"]
        rows.append([condition+" minus A0",f"{value['mean_difference']:.3f}",f"[{value['ci95'][0]:.3f}, {value['ci95'][1]:.3f}]"])
    table(doc,["Paired comparison","Mean delay difference ms","95 percent seed interval"],rows,[1.7,2.2,2.6])
    para(doc,"These intervals quantify variability across five seeds in one synthetic workload family. They do not capture topology variation, model mismatch or deployment uncertainty. Offline decisions are deterministic and cannot establish a language model benefit.")

    page(doc,"7 Reflection and memory findings")
    image(doc,"offline_reward.png","Figure 2. Ten-window rolling reward averaged over seeds. Reward compares each action with its same-window no-op counterfactual.")
    para(doc,f"Mean reward was {c['A1']['reward']:.4f} for A1 and {c['A2']['reward']:.4f} for A2. The rolling curve is not evidence of sustained improvement simply because it fluctuates above zero. The central comparison is the controller with memory against the same controller without memory, with goodput and total loss preserved.")
    para(doc,"Three mechanisms can explain the observed weakness. Cosine similarity over a small feature vector may fail to distinguish states near a feasibility boundary. Reward-based pruning favors previously successful episodes while dropping failures that could warn the controller. Reusing an action label from a related state does not guarantee that the action has the same effect under current path demand.")
    para(doc,"The safety gate prevents invalid reroutes from reaching the simulator, but vetoes still consume a decision opportunity. A possible next experiment would filter retrieved actions by present feasibility, retain informative negative examples, and compare reward-ranked, similarity-ranked and recent-window memories. Those changes require a fresh preregistered evaluation; their expected benefit is a hypothesis, not a result in this submission.")
    doc.add_heading("Meaning of self improvement",level=2)
    para(doc,"The implementation supports adaptation through context. It does not change model weights, train a new embedding or establish monotonic learning. A complete execution trace demonstrates that the loop works technically. Whether it improves control quality is an empirical question, and this offline result does not support a general improvement claim.")

    page(doc,"8 Live Gemini integration pilot")
    result_table(doc,gemini)
    g1=coverage(gemini,"A1");g2=coverage(gemini,"A2")
    para(doc,f"The paced pilot requested gemini-3.5-flash-lite. A1 received API decisions in {g1[0]} of {g1[1]} windows; A2 received API decisions in {g2[0]} of {g2[1]} windows. The remaining windows are explicit safe fallbacks. Every actual model identifier, duration, response status and returned usage field is stored in provider_calls.json. Network results remain simulated.")
    image(doc,"api_latency.png","Figure 3. API pilot mean modeled latency. Different pilot sizes and failure coverage prevent a fair cross-provider performance ranking.")
    para(doc,f"Gemini A1 mean modeled latency was {gc['A1']['latency_ms']:.3f} ms and A2 was {gc['A2']['latency_ms']:.3f} ms. A1 mean total loss was {gc['A1']['loss_pct']:.3f}% and A2 was {gc['A2']['loss_pct']:.3f}%. These values include all attempted operational windows. A successful model response can still be vetoed, and a valid no-op decision need not improve a congested network.")
    para(doc,"An earlier unpaced run hit HTTP 429 and is retained as an availability diagnostic in a separate directory. The corrected adapter spaces request starts, caps response wait time and opens a circuit breaker after consecutive failures. The paced run uses its own output directory; failed evidence has not been overwritten. Provider quotas depend on project and model [5], so request spacing does not guarantee availability or a free-tier entitlement.")
    successful=sum(r["status"]=="ok" for r in gemini["audit"])
    para(doc,f"The paced run completed {successful} successful provider requests. Its measured wall time was {gemini['summary']['wall_seconds']:.1f} seconds. Request spacing and two calls per control cycle mean that the live API loop is slower than a five-second real-time deadline. Simulation time advances independently; this pilot does not certify real-time control.")

    page(doc,"9 Free models and API availability")
    image(doc,"api_coverage.png","Figure 4. API decision coverage. Safe fallbacks are counted separately and remain in the performance results.")
    o1=coverage(orouter,"A1");o2=coverage(orouter,"A2")
    para(doc,f"OpenRouter's public catalog exposed {total} zero-price model IDs ending in :free. One bounded healthy-network decision was attempted for each ID. {passed} returned valid typed decisions and {total-passed} were unavailable or unsuitable in that check. Failure codes include rate limits, access restrictions and non-JSON or incomplete output. A healthy-network check tests access and schema handling; it does not measure congestion-control competence.")
    para(doc,f"The free-router diagnostic received API decisions in {o1[0]}/{o1[1]} A1 windows and {o2[0]}/{o2[1]} A2 windows. It exposed a practical limitation: a router can select models with different response behavior, including reasoning outputs that do not produce a usable JSON object within the configured token budget. Those requests became visible fallbacks. The actual underlying model and returned cost are recorded for every successful call.")
    para(doc,"The application blocks paid OpenRouter model IDs. Free routing is useful for a low-cost integration demo, but changing model identity and variable availability confound a controlled model-quality comparison. The diagnostic is therefore reported as operational evidence rather than a proof that one provider or memory strategy performs better.")
    para(doc,"A fixed-model follow-up should select one accessible free model, verify its output budget on a small case, and evaluate A1/A2 on multiple identical seeds with sufficient successful decision coverage. No reported literature score substitutes for this local experiment.")

    page(doc,"10 Engineering validation and reproducibility")
    para(doc,"Regression checks cover seeded trace equality, conservation of traffic volume, throughput and queue bounds, invalid targets, alternate overload, embedding retrieval, atomic memory persistence and explicit API failure provenance. The original graph, reward, counters, safety and validated result bundle tests remain part of the suite.")
    para(doc,"Dashboard checks verify loopback operation, saved comparisons, a working offline replay, individual decision navigation and mobile layout. The HTTP server serves a fixed route allowlist. Requests for .env, source configuration, repository metadata and traversal paths are rejected. Replay requires a same-origin request and accepts bounded integer parameters.")
    doc.add_heading("Corrections to the supplied project",level=2)
    for text in ["The original result writer defaulted outside selfx; its output now stays under output/testbed_results.", "A meter-only OpenFlow override omitted forwarding. The graph now preserves an observed output action and rejects a meter action when forwarding cannot be preserved.", "The linear testbed offered no alternate path. Rerouting now requires a tested egress mapping rather than trusting a model-generated port.", "Python test imports assumed a parent-level package. They now run from the requested folder.", "Gemini and free OpenRouter replace the previous single-provider dependency for current use. Provider failures are sanitized and visible."]:
        para(doc,text)
    doc.add_heading("Reproduction sequence",level=2)
    para(doc,"From selfx, run setup.ps1 on a new Windows machine, then run.ps1 test, run.ps1 experiment and run.ps1 dashboard. Existing saved API evidence can be reviewed without making new requests. New live runs require valid credentials and remaining provider quota. The README gives Python equivalents and exact output paths. Rebuilding report content uses scripts/build_artifacts.py; presentation authoring uses the Artifact Tool build module.")

    page(doc,"11 Limitations and next validation")
    para(doc,"The strongest supported conclusion is that SELF-X provides a traceable local control loop with real provider integration and bounded interventions. The offline memory variant did not consistently outperform its memory-free comparator. The API pilots have a small sample size and operational failures, which limit claims about model quality and memory benefit.")
    para(doc,"The fluid model uses constant demand within five-second windows, fixed path capacity and a simplified queue-delay proxy. It does not implement packet scheduling, TCP congestion control, real NetFlow collection, propagation variation, OpenFlow control-plane delay, link failures or real hardware. All reported path-performance measurements are simulation values. No external traffic dataset calibrates this model.")
    para(doc,"The numeric embedding is manually constructed. Cosine similarity does not ensure that a retrieved episode is causally relevant, and reward pruning can bias the memory. The reward weights encode a selected tradeoff among latency, goodput and loss; changing them can change the preferred action. A formal stability or SLA guarantee is not established.")
    para(doc,"The Windows evaluation cannot certify Linux Mininet, Open vSwitch, Ryu, iperf3 or root networking. The retained linear topology needs endpoint instrumentation for goodput and end-to-end loss; its aggregate switch counters count transmissions across hops. Its installation and sustained 600-second packet-emulation runs remain deployment work, separate from the accelerated simulation.")
    doc.add_heading("Recommended research sequence",level=2)
    para(doc,"First, add endpoint measurements and validate topology-specific action mappings in a dedicated Linux testbed. Next, compare memory retrieval variants with fixed models and adequate API coverage. Then replay an appropriately licensed captured workload and vary topology, link capacity and congestion patterns. Report matched baselines and uncertainty without dropping failures. Only those studies can support broader claims about real-network performance or generalization.")

    page(doc,"References and evidence files")
    for i,(title,url) in enumerate(SOURCES,1):
        para(doc,f"[{i}] {title} {url}")
    para(doc,"Web sources accessed 20 October 2026. Reported RouteNet prediction error belongs to its original study and is not reproduced or compared as an SDN control improvement here.")
    doc.add_heading("Local evidence",level=2)
    for text in ["output/experiments/offline: five-seed deterministic study, CSV metrics, JSONL trace and traffic hashes.", "output/experiments/gemini_paced: paced Gemini pilot and actual provider audit.", "output/experiments/gemini: unpaced quota diagnostic retained separately.", "output/experiments/openrouter: free-router integration diagnostic.", "output/provider_check.json and output/free_model_checks.json: authenticated connectivity and complete discovered free-model sweep.", "docs/ARCHITECTURE.md, docs/ARCHITECTURE.Rmd and docs/EXPERIMENTS.md: module flow and exact protocol."]:
        para(doc,text)
    FINAL.mkdir(parents=True,exist_ok=True)
    benchmark_data=OUT/'benchmark_chart_data.json'
    if benchmark_data.exists():
        from scripts.benchmark_report import append
        append(doc,json.loads(benchmark_data.read_text(encoding='utf-8')),page,table)
    path=FINAL/"SELF_X_Report.docx";doc.save(path)
    return path


def slide_content(runs,checks):
    # Native charts in the deck use these exact literal values.
    data={"title":"SELF-X", "subtitle":"SDN control with state embedding memory", "date":"20 October 2026",
          "sources":SOURCES,"runs":{name:value["summary"] for name,value in runs.items()},
          "coverage":{name:{c:coverage(value,c) for c in ("A1","A2")} for name,value in runs.items()},
          "free_models":{"checked":len(checks["models"]),"ok":sum(x["status"]=="ok" for x in checks["models"])}}
    if (OUT/'benchmark_chart_data.json').exists():
        data['benchmark']=json.loads((OUT/'benchmark_chart_data.json').read_text(encoding='utf-8'))
        data['runtime']=json.loads((OUT/'runtime_validation.json').read_text(encoding='utf-8'))
        if (OUT/'experiments/gemini_venv/summary.json').exists():
            fresh=load_run('gemini_venv');data['fresh_gemini']=fresh['summary']
            data['fresh_gemini_calls']={'ok':sum(r['status']=='ok' for r in fresh['audit']),'total':len(fresh['audit'])}
    (OUT/"deck_content.json").write_text(json.dumps(data,indent=2),encoding="utf-8")


def main():
    if (OUT/'benchmarks/v4/latency_gemini_final/summary.json').exists():
        from scripts.build_v4_artifacts import main as latest
        return latest()
    if (OUT/'benchmarks/v3/latency_gemini/summary.json').exists():
        from scripts.build_v3_artifacts import main as latest
        return latest()
    if (OUT/'benchmarks/v2/summary.json').exists():
        from scripts.benchmark_visuals import create
        create()
    runs={name:load_run(name) for name in ("offline","gemini_paced","openrouter")}
    checks=json.loads((OUT/"free_model_checks.json").read_text(encoding="utf-8"))
    figures(runs); path=make_report(runs,checks);slide_content(runs,checks)
    print(path)


if __name__=="__main__":main()
