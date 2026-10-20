"""Build the current editable report and chart data from saved v3 measurements."""
import json,csv
from pathlib import Path
from statistics import mean
from docx import Document
from docx.shared import Inches
from scripts.build_artifacts import style_document,para,table,page
from scripts.v3_visuals import create,ROOT,FOLDER,FIG,SHORT

def build_report(d):
    s=d['primary'];overall=d['overall'];live=d['live'];g=live['groups'];sensitivity=d['sensitivity'];doc=Document();style_document(doc)
    def p(text):para(doc,text)
    def picture(name,caption):
        doc.add_picture(str(FIG/(name+'.png')),width=Inches(6.65));doc.add_paragraph(caption,style='Caption')
    def fmt(v):return f'{v:.3f}'
    doc.add_paragraph('SELF X SDN control with adaptive episodic memory',style='Title')
    doc.add_paragraph('Project report | Measured revision 3 | 20 October 2026',style='Subtitle')
    doc.add_heading('Abstract',level=1)
    p('SELF-X studies a single network controller that retrieves its own earlier state-action-outcome episodes before selecting a bounded flow intervention. The local website exposes generated traffic, action validation, measured network outcomes and provider provenance. This revision improves memory retrieval and reduces provider work while retaining a deterministic action safety gate.')
    p(f"The primary evaluation contains {s['window_count']:,} simulated control windows across eight traffic scenarios, nine controllers and ten held-out seeds. Adaptive memory A4 preserves load magnitude, separates action values, retains diverse episodes and requires local support before overriding its guarded base policy. Mean modeled delay was {overall['A4']['latency_ms']:.3f} ms for A4 and {overall['A3']['latency_ms']:.3f} ms for previous guarded memory A3. A4 matched guard-only H1 in this suite; this is a regression reduction rather than evidence that memory beats the shared guard.")
    p(f"A separate constructed near-capacity diagnostic contains {sensitivity['window_count']:,} windows. It demonstrates adaptation after harmful throttling, with better reward and loss at the cost of higher modeled delay. It is a post-hoc mechanism test and is excluded from the primary average.")
    p(f"The fresh Gemini pilot used two seeds and 24 cycles per mode. Mean full-cycle wall time fell from {g['legacy']['control_elapsed_ms']/1000:.3f} seconds for two serial calls to {g['compact']['control_elapsed_ms']/1000:.3f} seconds for a compact single call and {g['hybrid']['control_elapsed_ms']/1000:.3f} seconds for the hybrid path. Hybrid reduction was {d['latency_reduction_pct']['hybrid']:.1f}%. Calls, cache hits, local bypasses and deadline misses remain separate. API requests are real; all network measurements are from a finite-buffer fluid simulation.")
    p(f"The corrected hybrid pilot reduced modeled delay but increased mean loss by {g['hybrid']['loss_pct']-g['legacy']['loss_pct']:.3f} percentage points relative to legacy. Hybrid missed the five-second target in {g['hybrid']['deadline_misses']} of 24 cycles. Lower average wall time therefore does not establish uniform control-quality improvement or a deadline guarantee.")
    p('The report includes methods, paired comparisons, timing decomposition, limitations and reproduction commands. Raw evidence, synthetic NetFlow-style records, PNG/SVG figures and their source hashes are under output/benchmarks/v3. No real-network or packet-testbed performance is claimed.')

    page(doc,'1 Project objective and control architecture')
    p('Repeated congestion can make a static primary route inefficient even when another path has spare capacity. SELF-X asks whether earlier local experiences can help a controller choose an appropriate action, and whether provider integration can operate within a practical control interval. The website is a local research and review dashboard, with an offline replay and submission downloads.')
    table(doc,['Stage','Operation','Audit evidence'],[['Observe','Read current demand and queue occupancy','Seeded traffic JSON and CSV'],['Retrieve','Match prior contexts and action values','Earlier episode cycle numbers'],['Decide','Use supported memory, a cached action or an API proposal','Decision origin and provider audit'],['Validate','Check exact target, route capacity and rate bounds','Allowed action or explicit fallback'],['Act and measure','Apply one five-second intervention','Conserved traffic volumes and model metrics'],['Reflect and store','Summarize measured outcome locally','Bounded persistent episode store']],[1.05,3.3,2.15])
    doc.add_heading('Local navigation',level=2)
    p('README.md is the entry point. docs/PROJECT_GUIDE.md explains the project goal and demo. docs/ARCHITECTURE.md and its R Markdown companion map modules and flow. docs/BENCHMARKS_V3.md defines the current evaluation. The dashboard compares the latest results and distinguishes API work from local reuse.')
    p('The portable model and the original Linux Mininet/Ryu adapter remain separate execution paths. This Windows submission validates the portable model; it does not certify Mininet, Open vSwitch, endpoint packet measurements or deployed OpenFlow rules.')

    page(doc,'2 Generated traffic and network model')
    p('Each logical path has a 10 Mbps bottleneck, a 0.5 Mbit finite queue and an 8 ms propagation component. One elephant flow begins each window on the primary path; background demand occupies both paths. Each action lasts one five-second window. Queues carry over within a controller and seed, and all controllers start with empty queues and memory.')
    table(doc,['Scenario','Generated behavior'],[[s['scenarios'][sc],text] for sc,text in zip(s['scenarios'],['Repeated light, burst, alternate-heavy and dual-congestion episodes','Six-window flash bursts followed by recovery','Alternate background already near its bottleneck','Alternate background rises halfway through the trace','Both paths remain persistently overloaded','Primary demand fits without intervention','Six seeded demand archetypes repeat exactly','Independent unpredictable demand changes'])],[2.35,4.15])
    p('The 9,600 unique primary workload windows are exported as traffic/simulated_netflow.csv. The file contains virtual start/end times, logical endpoints, elephant bytes offered and separate background rates. It is a NetFlow-style aggregate table, not captured packets or a wire-format IPFIX message. Original JSON traces retain SHA256 hashes.')
    p('For each path, final queue = min(buffer, max(0, previous queue + (admitted rate minus capacity) × duration)). Delivered volume is previous queue plus admitted volume minus overflow minus final queue. Total loss includes overflow and deliberate meter drops. Average queue occupancy is integrated over the window and converted into a modeled delay proxy, weighted by delivered volume across both paths.')

    page(doc,'3 Adaptive state embedding and memory')
    p('Earlier cosine retrieval could confuse states with similar proportions but different absolute load or feasibility. Reward-based pruning could discard informative failures. A3 added feasibility checks and balanced storage but still reused an action on optimistic local reward evidence. A4 makes the action-conditioned evidence gate more conservative.')
    table(doc,['Component','Previous A3','Current A4'],[['State match','Normalized cosine similarity','Capacity-scaled magnitudes and headroom'],['Features','Demand and queues','Demand, three capacity margins and queues'],['Action evidence','Weighted reward ranking','Separate local advantage estimate per action'],['Retention','At most 30 balanced episodes','At most 96 diverse episodes; four per action/context cell'],['Minimum support','Single similar positive case can influence reuse','Three matched prior observations'],['Uncertainty','Fixed positive margin','Weighted local spread plus a support margin'],['Feasibility','Validate reused action','Validate each candidate and rebuild current parameters']],[1.4,2.4,2.7])
    p('The context vector contains normalized primary background, alternate background and elephant demand; primary overload, reroute headroom and post-meter overload margins; and both queue occupancies. Matches must share the congestion/feasibility regime. Local distances retain magnitude rather than normalizing away load size. Cached vectors avoid repeated feature reconstruction.')
    p('At most eight matched observations per action contribute, with distance and recency weights. A4 chooses an alternative only when its lower evidence margin exceeds the base action estimate. Repeated negative evidence can select no-op. These local uncertainty margins are decision heuristics, not calibrated confidence intervals. The paired bootstrap intervals reported for experiments have a separate statistical role.')
    p('Only completed episodes with cycle < current cycle are eligible. Selection never evaluates candidate simulator outcomes. The reward remains the original no-op-relative weighted combination of modeled delay, goodput and total loss; that counterfactual is generated for measurement after selection. No model weights are trained or fine-tuned.')

    page(doc,'4 API delay changes and request provenance')
    table(doc,['Mode','Provider work','Local behavior'],[['Legacy','Decision and reflection are serial API calls','Full prompts and structured flow parameters'],['Compact','One short action-choice API call','Pooled HTTP session, local parameter construction and measured reflection'],['Hybrid','Compact call only when local reuse cannot serve','Healthy bypass, exact-context validated cache, supported episodic memory']],[1.1,2.7,2.7])
    p('The compact Gemini profile requests minimal thinking and a smaller output budget. The prompt supplies computed current primary/alternate totals and feasibility flags so the model need not repeat arithmetic. Prior examples include their own load context. Ports, flow matches and current rate limits are constructed locally and checked independently. The model name and actual returned version are recorded in each provider audit.')
    p('The cache is bounded to 64 entries with a 12-window expiry. Its exact signature includes demand, queues, path capacity, buffer size and observed route ports. On every hit, current flow parameters are rebuilt and validated. Cached actions are labeled api_cache; they are not counted as fresh model decisions. Local reflection uses the measured outcome and makes no API request.')
    p('Request pacing is retained to respect the configured provider rate. The adapter records pacing time separately from HTTP/network/service time. Full-cycle timing includes the decision, validation, outcome measurement, reflection and memory update. Work is not silently moved to an unmeasured asynchronous stage.')
    p('An initial pilot on seeds 1770 and 1771 exposed incorrect arithmetic in compact responses. It is preserved as latency_gemini_initial. Explicit current-state totals and flags were then added, and the reported fresh pilot uses seeds 1790 and 1791. This is a small integration and system-latency evaluation, not a broad LLM-quality study.')

    page(doc,'5 Evaluation design and comparisons')
    table(doc,['Dataset','Seeds and windows','Purpose'],[['Primary held-out suite','1760–1769; 8 × 10 × 120 × 9 = 86,400','Compare controllers on matched generated traffic'],['Memory development','1740 and 1741','Select retrieval rules and correct implementation tests'],['Fresh Gemini timing pilot','1790 and 1791; 12 × 2 × 3 = 72','Compare two-call, compact and hybrid paths'],['Near-capacity sensitivity','1780–1789; 4 × 10 × 120 × 4 = 19,200','Post-hoc mechanism diagnostic; excluded from primary average'],['Free model checks',f"{d['free_models']['checked']} discovered zero-price model IDs",'One bounded healthy-network schema check per model']],[1.5,2.5,2.5])
    table(doc,['ID','Controller role'],[[c,s['controllers'][c]] for c in s['controllers']],[.7,5.8])
    p('The primary seeds are distinct from development seeds. Every controller receives identical demand and starts from zero queues. H1 isolates the base guard without episodic memory. A3 is the previous guarded-memory implementation. O1 evaluates current candidate simulator outcomes and is a privileged myopic reference; it is neither a deployment policy nor a long-horizon upper bound.')
    p('Seed means are the unit of inference. Paired seed differences use a fixed-RNG 4,000-draw percentile bootstrap and 95% intervals. Comparisons are exploratory without multiplicity correction. P95 network delay is the P95 of modeled window means, not packet RTT. The two-seed live pilot reports descriptive timing; its windows are not treated as independent seed replications.')

    page(doc,'6 Primary held out network results')
    picture('scenario_latency_matrix','Figure 1. Mean modeled delay per scenario and controller on the primary held-out suite.')
    table(doc,['Scenario','A4 minus A3 ms','Paired 95% interval'],[[SHORT[sc],fmt(v['latency_ms']['mean_difference']),f"[{v['latency_ms']['ci95'][0]:.3f}, {v['latency_ms']['ci95'][1]:.3f}]"] for sc in s['scenarios'] for v in [s['paired_comparisons'][sc]['A4_minus_A3']]],[1.6,2,2.9])
    p('A4 reduced A3 delay regressions on mixed, flash and volatile demand in this run. Several scenarios were unchanged. Primary conclusions retain all windows and all ten seeds, including infeasible proposals and cold-start periods.')

    page(doc,'7 Memory effects and the guard ablation')
    picture('paired_memory_effect','Figure 2. A4 minus previous guarded memory A3, with paired seed uncertainty.')
    p(f"Across equally weighted primary scenarios, A4 minus A3 delay was {d['overall_paired']['A3']['latency_ms']['mean_difference']:.3f} ms. The paired seed interval was [{d['overall_paired']['A3']['latency_ms']['ci95'][0]:.3f}, {d['overall_paired']['A3']['latency_ms']['ci95'][1]:.3f}] ms. Mean total loss changed by {d['overall_paired']['A3']['loss_pct']['mean_difference']:.3f} percentage points.")
    p(f"A4 and H1 produced the same installed actions and network metrics across this primary suite. A4 retrieved supported evidence in {overall['A4']['retrieval_windows']:,} of 9,600 windows, but the conservative gate did not choose a different action from H1. This revision prevents observed A3 regressions; it does not establish that adding memory improves the guard-only policy on ordinary traffic.")
    p('This distinction matters because a memory controller can appear better after a heuristic is added even when memory contributes no additional performance. The separate H1 comparison prevents that attribution error. The targeted diagnostic in Section 9 tests a specific negative-experience mechanism rather than altering the primary evaluation average.')

    page(doc,'8 Delay goodput loss and action feasibility')
    table(doc,['ID','Delay ms','Goodput Mbps','Total loss %','Vetoes'],[[c,fmt(v['latency_ms']),fmt(v['throughput_kbps']/1000),fmt(v['loss_pct']),str(v['vetoes'])] for c,v in overall.items()],[.6,1.4,1.6,1.5,1.4])
    picture('metric_tradeoffs','Figure 3. Primary scenario averages across delay, loss and delivered goodput. A4 overlaps H1.')
    p('Lower modeled delay alone is insufficient. Deliberate meter drops reduce queue pressure while discarding offered traffic. All loss metrics therefore include meter drops, and delivered goodput is reported alongside delay. A4 had zero safety vetoes; feasibility and performance are separate properties.')

    page(doc,'9 Targeted near capacity memory diagnostic')
    picture('memory_sensitivity','Figure 4. Constructed sensitivity experiment outside the primary average. All cold-start windows are retained.')
    first=next(iter(sensitivity['groups'].values()));h=first['H1']['metrics'];a=first['A4']['metrics']
    p('This post-hoc experiment was constructed to exercise a specific failure mechanism: a load only slightly above capacity can trigger a fixed 20% meter even when its goodput cost exceeds the queue-delay benefit. Four overload levels are evaluated on ten matched seeds each, with the alternate path unavailable for safe rerouting.')
    p(f"At 0.0001 Mbps excess, H1 mean total loss was {h['loss_pct']['mean']:.3f}% and A4 was {a['loss_pct']['mean']:.3f}%. Mean reward changed from {h['reward']['mean']:.6f} to {a['reward']['mean']:.6f}. A4 delay increased from {h['latency_ms']['mean']:.3f} to {a['latency_ms']['mean']:.3f} ms. This is a measured tradeoff, not an improvement in every metric.")
    table(doc,['Excess Mbps','A4 minus H1 reward','A4 minus H1 loss pp'],[[str(sensitivity['levels_mbps'][level]),fmt(sensitivity['paired_comparisons'][level]['A4_minus_H1']['reward']['mean_difference']),fmt(sensitivity['paired_comparisons'][level]['A4_minus_H1']['loss_pct']['mean_difference'])] for level in sensitivity['levels_mbps']],[1.5,2.6,2.4])
    p('These results support the local negative-experience mechanism under a deliberately constructed workload. They do not establish generalization to captured traffic, a new topology or a real network.')

    page(doc,'10 Learning from completed harmful actions')
    picture('memory_learning','Figure 5. Cumulative reward and action selection after cold start in the 0.0001 Mbps sensitivity diagnostic.')
    p('A4 first observes the guarded rate-limit outcome. After at least three matched completed observations, repeatedly negative local action evidence can reject further throttling and select no-op. The queue and demand features keep capacity-boundary cases separate from strongly congested cases. The action validator still applies to each selected intervention.')
    p('The plotted cumulative average is descriptive. It is not a proof of monotonic learning, global stability or neural training. A4 records bounded experiences and adapts retrieval decisions; its language-model weights do not change. The privileged O1 reference has information unavailable to the policy and is displayed only as a contextual diagnostic.')
    p('The learning test uses the same measured outcome and reward definitions as the primary suite. It includes the initial harmful decisions rather than discarding them as warm-up. Raw cycles and their retrieved episode indices are available for independent review.')

    page(doc,'11 Measured Gemini API and full cycle delay')
    picture('api_latency_compare','Figure 6. Mean timing decomposition and P95 full-cycle wall time from the fresh Gemini pilot.')
    table(doc,['Mode','Mean cycle s','P95 cycle s','API calls','Misses out of 24'],[[mode,fmt(g[mode]['control_elapsed_ms']/1000),fmt(g[mode]['p95_control_ms']/1000),str(g[mode]['api_calls']),str(g[mode]['deadline_misses'])] for mode in ('legacy','compact','hybrid')],[1.1,1.4,1.4,1.15,1.45])
    p(f"The compact path reduced mean full-cycle time by {d['latency_reduction_pct']['compact']:.1f}% relative to the two-call path. Hybrid reduced it by {d['latency_reduction_pct']['hybrid']:.1f}%. Removing the reflection request, reducing prompt/output size and using a pooled connection change provider work; local cache and bypass reduce the number of new calls.")
    p('Transport time includes HTTP/network and service latency; the experiment does not isolate provider-server execution or TLS savings. Pacing remains in the full-cycle total. A zero-miss outcome in a 24-window pilot would not establish a production deadline guarantee, and the adapter retains explicit fallback behavior when providers fail.')

    page(doc,'12 API decision origins and network quality')
    picture('api_call_origins','Figure 7. Actual provider requests and decision provenance per mode.')
    table(doc,['Mode','Delay ms','Goodput Mbps','Total loss %','Reward'],[[m,fmt(g[m]['latency_ms']),fmt(g[m]['throughput_kbps']/1000),fmt(g[m]['loss_pct']),fmt(g[m]['reward'])] for m in ('legacy','compact','hybrid')],[1.05,1.3,1.55,1.4,1.2])
    p('All API requests, successful decisions, cache-served windows, healthy bypasses and fallbacks remain identifiable in metrics.csv and trace.jsonl. Faster execution must be judged together with control quality, not on timing alone. A cached action is still validated against the present state, and a new model response is not automatically evidence that a good action was chosen.')
    p('The initial compact arithmetic failure is preserved rather than overwritten. Its poorer network outcomes motivated explicit locally computed current-state totals and feasibility flags. The new pilot evaluates that corrected prompt on fresh traffic seeds. This small result supports technical integration and system-latency comparisons; it is not a cross-provider or cross-model quality ranking.')

    page(doc,'13 Cycle traces and free model availability')
    picture('api_timeline','Figure 8. Full-cycle timing by traffic window, with a five-second target.')
    p('Mode order is reversed on the second seed to reduce simple time-order bias. Provider variation, cache warm-up and pacing still affect results. The hybrid benefit depends on repeated exact contexts; workloads with few repeats will have fewer cache hits and may approach the compact path rather than the measured hybrid average.')
    p(f"The current compact OpenRouter sweep checked all {d['free_models']['checked']} discovered zero-price :free model IDs; {d['free_models']['ok']} returned valid healthy-network action choices. The remaining models were unavailable or unsuitable under the configured output budget. Every result and sanitized provider audit is saved in free_model_checks.json. The atlas lists elapsed request times for each ID.")
    p('These availability checks are separate from the fixed Gemini timing pilot. Free-model quota and availability can change, and a simple healthy-network schema response does not measure congestion-control skill. The application blocks paid OpenRouter model IDs.')

    page(doc,'14 Engineering validation and practical limits')
    runtime=json.loads((ROOT/'output/runtime_validation.json').read_text())
    p('The full portable workflow ran with this folder’s isolated .venv interpreter. Automated checks cover deterministic traffic, traffic conservation, action feasibility, episode chronology, persistence, memory support, context magnitude, negative-experience learning, exact cache expiry, state changes, healthy bypass and explicit provider failures.')
    p(runtime['checks'][0]['output_tail'].strip().splitlines()[-1])
    picture('offline_control_overhead','Figure 9. Measured local control-loop computation mean and P95. Final batched persistence is excluded.')
    p(f"The primary suite completed in {s['wall_seconds']:.2f} seconds on this computer. A4 mean computation was {overall['A4']['control_elapsed_ms']:.3f} ms, with P95 {overall['A4']['decision_p95_ms']:.3f} ms. These machine-dependent values are separate from real API waiting.")
    p('All network values use a fixed-topology fluid model. Packet scheduling, TCP behavior, real flow export, hardware timing and control-plane delay are not modeled. The Linux adapter needs separate endpoint instrumentation and packet-testbed validation. The memory’s local evidence margins are not formal stability, fairness or SLA guarantees. Workload diversity and real deployment remain future validation tasks.')

    page(doc,'15 Reproduction sources and submission navigation')
    p('From the selfx folder, run setup.ps1 on a new machine to create .venv and install dependencies. The .env file stores credentials separately and remains excluded from the submission archive. Existing saved results can be reviewed without new API calls.')
    for command in [r'.\run.ps1 all',r'.\run.ps1 latency',r'.\.venv\Scripts\python.exe -m eval.run_memory_sensitivity',r'.\.venv\Scripts\python.exe -m scripts.export_simulated_netflow',r'.\run.ps1 dashboard -Port 8766']:
        p(command)
    p('The all command runs tests, a complete portable experiment and the primary held-out benchmark. The latency command makes real provider requests and consumes available quota. The sensitivity command runs the separate mechanism diagnostic. scripts/build_v3_artifacts.py regenerates the report and figures from saved results; the editable PowerPoint builder is scripts/build_v3_presentation.mjs.')
    doc.add_heading('Evidence and editable outputs',level=2)
    p('output/benchmarks/v3 contains raw metrics, demand traces, traffic hashes, generated NetFlow-style records, API audits, the sensitivity experiment, free-model checks and the dedicated figures folder. output/submission contains the editable DOCX and PPTX, report and slide PDFs, the graph atlas, verification and the complete ZIP. Earlier results are preserved under output/benchmarks/v2 and output/history/v2.')
    doc.add_heading('Selected primary sources',level=2)
    for text in ['Shinn et al. Reflexion (2023): https://arxiv.org/abs/2303.11366',
        'Google Gemini 3.5 Flash-Lite model: https://ai.google.dev/gemini-api/docs/models/gemini-3.5-flash-lite',
        'Google thinking levels: https://ai.google.dev/gemini-api/docs/generate-content/whats-new-gemini-3.5',
        'OpenRouter free router: https://openrouter.ai/docs/guides/routing/routers/free-router',
        'RFC 2992 ECMP analysis: https://www.rfc-editor.org/rfc/rfc2992',
        'RFC 7011 IPFIX protocol: https://www.rfc-editor.org/rfc/rfc7011']:
        p(text)
    p('These sources motivate the method and interfaces. Published results are not substituted for locally measured controller outcomes. Sources checked on 20 October 2026; model and free-tier availability can change.')
    path=ROOT/'output/submission/SELF_X_Report.docx';doc.save(path);return path

def main():
    d=create();print(build_report(d));return d

if __name__=='__main__':main()
