"""Enhanced benchmark appendix for the editable report."""
import json
from pathlib import Path
from docx.shared import Inches

ROOT=Path(__file__).resolve().parents[1]


def append(doc,data,page,table):
    def paragraph(text):doc.add_paragraph(text)
    def figure(name,caption,width=6.5):
        doc.add_picture(str(ROOT/'output/figures/benchmarks'/(name+'.png')),width=Inches(width))
        doc.add_paragraph(caption,style='Caption')
    summary=json.loads((ROOT/'output/benchmarks/v2/summary.json').read_text(encoding='utf-8'))
    overall=data['overall'];order=list(data['controllers'])
    page(doc,'Appendix A Expanded benchmark protocol')
    paragraph(f"The enhanced suite evaluates {data['count']:,} windows across six demand scenarios, eight controllers and ten paired seeds. Each condition and seed runs 120 five-second windows. The measured suite runtime was {data['wall_seconds']:.2f} seconds in the isolated project .venv. All path-performance values are analytical simulation outputs.")
    table(doc,['ID','Controller and information available'],[[c,data['controllers'][c]] for c in order],[.65,5.85])
    paragraph('Mixed demand repeats congestion regimes. Flash-crowd demand creates bursts on the primary while an alternate has headroom. A busy-alternate scenario limits rerouting options. Distribution shift changes alternate background demand halfway through the run. Persistent overload tests behavior when both paths are constrained. Light load is a negative control.')
    paragraph('B2 uses a stable SHA256 hash of each synthetic flow ID over two equal key-space regions, inspired by the hash-threshold idea described in RFC 2992. The common safety gate still applies, so this is a guarded ECMP-style abstraction rather than a complete ECMP implementation. Synthetic records are not IPFIX messages or a measured NetFlow export; RFC 7011 defines the actual export protocol.')
    paragraph('O1 evaluates feasible no-op, reroute and 20% meter candidates with full knowledge of the simulator, selecting the highest immediate reward. This is a privileged myopic reference. It is not a deployable learned controller or a guaranteed long-horizon upper bound. Its reward objective also combines several metrics, so it need not minimize delay alone.')
    paragraph('Primary sources: C. Hopps, RFC 2992, https://www.rfc-editor.org/rfc/rfc2992 ; B. Claise et al., RFC 7011, https://www.rfc-editor.org/rfc/rfc7011 . Accessed 20 October 2026.')

    page(doc,'Appendix B Scenario results and uncertainty')
    figure('scenario_latency_matrix','Figure A1. Mean modeled window delay in milliseconds. Every matrix cell averages 1,200 windows across ten paired seeds.')
    figure('scenario_latency_intervals','Figure A2. Scenario-specific delay means and 95% bootstrap intervals across seed means. Controller IDs are defined in Appendix A.')
    paragraph('Bootstrap intervals use 4,000 fixed-seed resamples of the ten seed-level estimates. Pairwise differences resample paired seed differences. Windows are not treated as independent replicates. These exploratory comparisons have no multiplicity adjustment; an interval excluding zero is not a universal performance guarantee.')

    page(doc,'Appendix C Guarded memory and its ablation')
    figure('paired_memory_effect','Figure A3. Negative paired differences favor A3. The H1 comparison isolates memory after the shared guard.')
    paragraph(f"Original A2 produced {overall['A2']['vetoes']} rejected proposals in the expanded suite. A3 produced {overall['A3']['vetoes']}. A3 rebuilds action parameters for current demand and filters retrieved actions through the same deterministic feasibility rules used by the final gate. Its memory balances recent episodes with positive and negative reward examples and stays bounded at 30 records.")
    paragraph('H1 shares the guard that avoids intervention when offered primary load fits capacity. A3 uses the same base policy plus the enhanced memory. This extra comparator prevents a heuristic change from being reported as a memory benefit. A3 never queries simulator outcomes or future demand to select its action; only the privileged O1 reference does so.')
    mixed=summary['paired_comparisons']['mixed']['A3_minus_H1']['latency_ms']
    paragraph(f"On mixed demand, A3 minus H1 mean delay was {mixed['mean_difference']:.3f} ms, with paired 95% interval [{mixed['ci95'][0]:.3f}, {mixed['ci95'][1]:.3f}]. A3 was slower in that comparison. In the flash-crowd case the point estimate was also slower, and the other four scenario means were equal. Removing infeasible action reuse therefore demonstrates an engineering improvement in proposal validity, while a general memory-performance advantage remains unsupported.")
    paragraph('Balanced retention and feasibility filtering were introduced before the expanded run. The original A2 remains in the suite so unfavorable earlier behavior is still visible. More workloads, independently selected seeds and model-driven A3 pilots would be needed to establish broader findings.')

    page(doc,'Appendix D Distributions and metric tradeoffs')
    figure('distribution_ecdf','Figure A4. Empirical distributions pooled over equally sampled scenarios and seeds. Observations are modeled window means, not packet RTTs.')
    figure('metric_tradeoffs','Figure A5. Equal-weight scenario and seed averages. Interpret lower delay together with total loss and delivered goodput.')
    paragraph('The reported P95 statistic is the 95th percentile of modeled mean-window delays within each seed, averaged over seeds where summarized. It must not be described as packet-level P95 latency. Total loss includes deliberate meter drops, avoiding a misleading improvement from removing offered traffic from the denominator.')

    page(doc,'Appendix E Actions queues and control overhead')
    figure('actions_and_vetoes','Figure A6. Rejected proposals and installed action proportions across all scenarios.')
    figure('queue_shift_heatmap','Figure A7. One illustrative seed during distribution shift; alternate background demand increases at the dashed line.')
    paragraph('Queue occupancy is a model state in Mbit. The selected seed illustrates behavior and is not a substitute for aggregate paired estimates. Each controller starts from zero queues; every scenario and seed resets memory. Original reward-pruned A2 and balanced A3 receive only earlier episodes from their own condition.')

    page(doc,'Appendix F Runtime and reproducible execution')
    figure('offline_control_overhead','Figure A8. Mean and P95 control-loop computation time on this computer. Final batched memory persistence is excluded from these per-window timings.')
    table(doc,['Controller','Mean delay ms','Goodput Mbps','Mean loss %'],[[c,f"{overall[c]['latency_ms']:.3f}",f"{overall[c]['throughput_kbps']/1000:.3f}",f"{overall[c]['loss_pct']:.3f}"] for c in order],[1.05,1.75,1.75,1.95])
    paragraph('The .env file stores provider credentials and settings. The .venv directory is an isolated Python environment; these serve different purposes. Run .\\run.ps1 all to execute tests, the portable experiment and the expanded benchmark in .venv. Run .\\run.ps1 dashboard to start the local website. New machines use setup.ps1 first.')
    fresh=ROOT/'output/experiments/gemini_venv/provider_calls.json'
    if fresh.exists():
        audit=json.loads(fresh.read_text(encoding='utf-8'))
        paragraph(f"A fresh .venv Gemini integration pilot completed {sum(r['status']=='ok' for r in audit)} of {len(audit)} provider requests successfully. This is separate from the previously completed 288-call paced pilot. The additional pilot has one seed and 12 windows per controller, so it is an integration check rather than a strong model-quality benchmark. Current OpenRouter availability is recorded by the new .venv catalog sweep.")
